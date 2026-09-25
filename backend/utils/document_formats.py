import re
from io import BytesIO
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Inches

from fpdf import FPDF


def sanitize_text(text: str) -> str:
    """
    Clean text while preserving normal punctuation and structure.
    """

    if not text:
        return ""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\u2022": "-",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def text_to_terms(terms: str) -> List[str]:
    """
    Convert semicolon-separated terms into a list.
    """

    if not terms:
        return []

    return [
        item.strip()
        for item in terms.split(";")
        if item.strip()
    ]


def create_txt(text: str) -> bytes:
    """
    Generate a UTF-8 TXT file.
    """

    clean_text = sanitize_text(text)

    return clean_text.encode("utf-8")


def _is_heading(line: str) -> bool:
    """
    Determine whether a line looks like a document heading.
    """

    stripped = line.strip()

    if not stripped:
        return False

    if len(stripped) > 100:
        return False

    if stripped.startswith("#"):
        return True

    if re.match(r"^\d+[\.\)]\s+", stripped):
        return True

    uppercase_count = sum(
        1 for char in stripped if char.isupper()
    )

    letters_count = sum(
        1 for char in stripped if char.isalpha()
    )

    if letters_count > 4:
        return (
            uppercase_count / letters_count >= 0.75
        )

    return False


def create_docx(
    text: str,
    document_type: str,
    parties: str = "",
    effective_date: str = "",
    terms: str = ""
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Default font
    styles = document.styles

    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        document_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # Metadata
    if parties:
        p = document.add_paragraph()
        r = p.add_run("Parties: ")
        r.bold = True
        p.add_run(parties)

    if effective_date:
        p = document.add_paragraph()
        r = p.add_run("Effective Date: ")
        r.bold = True
        p.add_run(effective_date)

    document.add_paragraph()

    # Main content
    lines = sanitize_text(text).splitlines()

    for line in lines:

        stripped = line.strip()

        if not stripped:
            document.add_paragraph()
            continue

        if stripped.startswith("#"):
            stripped = stripped.lstrip("#").strip()

        if _is_heading(stripped):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(stripped)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            paragraph.space_before = Pt(8)
            paragraph.space_after = Pt(4)

        elif stripped.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                stripped[2:].strip()
            )

        else:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.15

            paragraph.add_run(stripped)

    # Terms table
    term_list = text_to_terms(terms)

    if term_list:

        document.add_paragraph()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "USER-PROVIDED KEY TERMS"
        )

        heading_run.bold = True
        heading_run.font.size = Pt(12)

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        header_cells = table.rows[0].cells

        header_cells[0].text = "No."
        header_cells[1].text = "Term"

        for index, term in enumerate(
            term_list,
            start=1
        ):
            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    # Footer
    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_paragraph.add_run(
        "Generated with LegalEase | AI-assisted legal drafting"
    )

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(self, document_type: str):
        super().__init__()

        self.document_type = document_type

        self.set_auto_page_break(
            auto=True,
            margin=20
        )

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            8,
            "LegalEase",
            align="C"
        )

        self.ln(5)

        self.set_draw_color(
            120,
            120,
            120
        )

        self.line(
            10,
            self.get_y(),
            200,
            self.get_y()
        )

        self.ln(5)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            10,
            "Generated with LegalEase | AI-assisted legal drafting",
            align="C"
        )


def create_pdf(
    text: str,
    document_type: str
) -> bytes:

    pdf = LegalEasePDF(
        document_type=document_type
    )

    pdf.set_margins(
        left=15,
        top=15,
        right=15
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        15
    )

    pdf.cell(
        0,
        10,
        sanitize_text(document_type).upper(),
        align="C"
    )

    pdf.ln(12)

    pdf.set_font(
        "Helvetica",
        "",
        10
    )

    clean_text = sanitize_text(text)

    for raw_line in clean_text.splitlines():

        line = raw_line.strip()

        if not line:
            pdf.ln(4)
            continue

        if line.startswith("#"):
            line = line.lstrip("#").strip()

        if _is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                line
            )

            pdf.ln(1)

            pdf.set_font(
                "Helvetica",
                "",
                10
            )

        elif line.startswith("- "):

            pdf.multi_cell(
                0,
                6,
                "- " + line[2:].strip()
            )

        else:

            pdf.multi_cell(
                0,
                6,
                line
            )

        pdf.ln(1)

    # fpdf2 returns bytearray in some versions.
    result = pdf.output()

    if isinstance(result, bytearray):
        return bytes(result)

    if isinstance(result, str):
        return result.encode("latin-1")

    return bytes(result)